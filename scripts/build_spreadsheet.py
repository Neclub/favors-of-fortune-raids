"""Build a workbook with one raid per sheet."""

import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
RAIDS = ROOT / "data" / "raids.json"
OUT = ROOT / "Info" / "FoF_Raid_Spells.xlsx"
FOLDERS = [
    Path(r"C:\Users\120ch\AppData\Local\Temp\fof-posts\searches"),
    Path(r"C:\Users\120ch\AppData\Local\Temp\fof-posts\spellpages"),
]
SLOT_RE = re.compile(r"Slot\s+\d+:\s*</b>\s*</td>\s*<td>\s*([^<]+)", re.I)
ROW_RE = re.compile(r"<b>\s*([^<]+?)\s*</b>\s*</td>\s*<td[^>]*>(.*?)</td>", re.I | re.S)
COUNTER_RE = re.compile(r"Increase ([A-Za-z]+) Counter by (\d+)", re.I)
DETAIL_ORDER = [
    "In Game Description",
    "Cast by You",
    "Cast by Other",
    "Land on You",
    "Land on Other",
    "Wear off",
    "Mana",
    "Endurance",
    "Casting Time",
    "Recast Time",
    "Duration",
    "Skill",
    "Target",
    "Range",
    "AE Range",
    "Max Targets",
    "Spell Type",
    "Dispellable",
    "Focusable",
    "Interruptible",
    "Restriction",
    "Push Back",
    "Push Up",
    "Viral",
    "Hate Override",
    "Spell Bar Lockout",
    "Stacking",
    "Spell Line",
    "Timer ID",
    "Category",
]
HEADERS = ["Spell", "Spell ID", "Spell information", "Counters", "Duration", "Land on You", "Wear off"]
SHEET_NAMES = {
    "siren": "Siren Sisters",
    "dipuc": "Dipuc",
    "love": "Scorn",
    "moruuk": "Moruuk and Sorshkah",
    "taurlton": "Taurlton",
    "vazaelle": "Fractured Vazaelle",
    "chronicles": "Chronicles of War",
    "eriak": "Eriak Zek",
    "xev": "Xev Tests Fate",
}


def flatten(raw):
    return "\n".join(line.strip() for line in raw.splitlines() if not re.fullmatch(r"\s*\d+\s*", line))


def clean(text):
    return re.sub(r"\s+", " ", text).strip()


def labeled(html, label):
    match = re.search(rf"{re.escape(label)}:</b>\s*</td>\s*<td>\s*([^<]*)", html, re.I)
    return clean(match.group(1)) if match else ""


def cell_text(fragment):
    text = re.sub(r"<br\s*/?>", " ", fragment, flags=re.I)
    text = re.sub(r"<[^>]+>", "", text)
    text = (
        text.replace("&nbsp;", " ")
        .replace("&amp;", "&")
        .replace("&#39;", "'")
        .replace("&quot;", '"')
        .replace("&lt;", "<")
        .replace("&gt;", ">")
    )
    return clean(text)


def spell_details(html):
    grouped = {}
    for label, raw in ROW_RE.findall(html):
        label = clean(label).rstrip(":").strip()
        if not label or label.lower() == "neclub" or re.fullmatch(r"Slot\s+\d+", label, re.I):
            continue
        if label.lower().startswith("drop down menu"):
            label = "Category"
        value = cell_text(raw)
        if not value:
            continue
        values = grouped.setdefault(label, [])
        if value not in values:
            values.append(value)
    details = []
    used = set()
    for label in DETAIL_ORDER:
        if label not in grouped:
            continue
        details.extend({"label": label, "value": value} for value in grouped[label])
        used.add(label)
    for label, values in grouped.items():
        if label in used:
            continue
        details.extend({"label": label, "value": value} for value in values)
    return details


def load_pages():
    pages = {}
    for folder in FOLDERS:
        if not folder.exists():
            continue
        for path in folder.glob("*.html"):
            html = flatten(path.read_text(encoding="utf-8", errors="replace"))
            if "Casting Time:" not in html and "Slot" not in html:
                continue
            spell_id = re.search(r"history\.php\?id=(\d+)", html) or re.search(r"spells\.php\?id=(\d+)", html)
            if not spell_id:
                continue
            slots = [clean(text) for text in SLOT_RE.findall(html)]
            pages[spell_id.group(1)] = {
                "slots": slots,
                "details": spell_details(html),
                "duration": labeled(html, "Duration"),
                "land": labeled(html, "Land on You"),
                "wear": labeled(html, "Wear off"),
            }
    return pages


def spell_fields(page):
    counters = []
    for text in page["slots"]:
        counter = COUNTER_RE.search(text)
        if counter:
            counters.append(f"{counter.group(1).title()} {counter.group(2)}")
    return {
        "info": "\n".join(page["slots"]),
        "counters": ", ".join(counters),
        "duration": page["duration"],
        "land": page["land"],
        "wear": page["wear"],
    }


def raid_spells(raid, spells):
    """Spell names the website links for this raid, in first-seen order."""
    ordered = []
    seen = set()
    names = sorted(spells, key=len, reverse=True)
    linkable = [name for name in names if " " in name or len(name) > 8]
    word = re.compile(r"[A-Za-z0-9'’]")

    def add(name):
        if name in spells and name not in seen:
            seen.add(name)
            ordered.append(name)

    def linked_names(text):
        if not text:
            return
        if text in spells:
            add(text)
            return
        used = [False] * len(text)
        hits = []
        for name in linkable:
            start = 0
            while True:
                index = text.find(name, start)
                if index < 0:
                    break
                end = index + len(name)
                before = index == 0 or not word.match(text[index - 1])
                after = end == len(text) or not word.match(text[end])
                if before and after and not any(used[index:end]):
                    for cursor in range(index, end):
                        used[cursor] = True
                    hits.append((index, name))
                start = index + 1
        for _, name in sorted(hits):
            add(name)

    def walk(bullets):
        for bullet in bullets:
            linked_names(bullet.get("text") or "")
            if bullet.get("spell"):
                add(bullet["spell"])
            walk(bullet.get("children") or [])

    for stage in raid["stages"]:
        for block in stage["blocks"]:
            linked_names(block.get("title") or "")
            for name in block.get("spellNames") or []:
                add(name)
            for paragraph in block.get("prose") or []:
                linked_names(paragraph)
            walk(block.get("bullets") or [])
    for row in raid.get("summary") or []:
        add(row["name"])
        linked_names(row.get("effect") or "")
    return ordered


def detail_value(match, label):
    values = [row["value"] for row in match.get("details") or [] if row.get("label") == label]
    return values[0] if values else ""


def main():
    data = json.loads(RAIDS.read_text(encoding="utf-8"))
    pages = load_pages()
    book = Workbook()
    book.remove(book.active)
    header_font = Font(bold=True, color="1B212B")
    header_fill = PatternFill("solid", fgColor="E2C57A")
    link_font = Font(color="0563C1", underline="single")
    band_fills = [
        PatternFill("solid", fgColor="F7F4EC"),
        PatternFill("solid", fgColor="E7EEF8"),
    ]
    wrap = Alignment(wrap_text=True, vertical="top")

    for tier in data["tiers"]:
        for raid in tier["raids"]:
            sheet = book.create_sheet(SHEET_NAMES[raid["id"]][:31])
            sheet.append(HEADERS)
            for cell in sheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="center")
            for name in raid_spells(raid, data["spells"]):
                matches = data["spells"].get(name) or []
                if not matches:
                    continue
                for match in matches:
                    page = pages.get(match["id"], {"slots": [], "duration": "", "land": "", "wear": ""})
                    slots = match.get("slots") or page.get("slots") or []
                    if not slots and match.get("effect"):
                        slots = [match["effect"]]
                    fields = spell_fields({"slots": slots, "duration": "", "land": "", "wear": ""})
                    fields["duration"] = detail_value(match, "Duration") or page.get("duration", "")
                    fields["land"] = detail_value(match, "Land on You") or page.get("land", "")
                    fields["wear"] = detail_value(match, "Wear off") or page.get("wear", "")
                    sheet.append([
                        name,
                        int(match["id"]),
                        fields["info"],
                        fields["counters"],
                        fields["duration"],
                        fields["land"],
                        fields["wear"],
                    ])
                    id_cell = sheet.cell(sheet.max_row, 2)
                    id_cell.hyperlink = f"https://spells.eqresource.com/spells.php?id={match['id']}&source=beta"
                    id_cell.font = link_font
            widths = [28, 12, 62, 18, 14, 55, 40]
            for column, width in enumerate(widths, start=1):
                sheet.column_dimensions[get_column_letter(column)].width = width
            previous = None
            band = 0
            for row in sheet.iter_rows(min_row=2):
                if row[0].value != previous:
                    if previous is not None:
                        band += 1
                    previous = row[0].value
                fill = band_fills[band % 2]
                needed = 1
                for cell, width in zip(row, widths):
                    cell.fill = fill
                    cell.alignment = wrap
                    text = "" if cell.value is None else str(cell.value)
                    lines = 0
                    for line in text.splitlines() or [""]:
                        lines += max(1, -(-len(line) // max(width - 1, 1)))
                    needed = max(needed, lines)
                sheet.row_dimensions[row[0].row].height = needed * 16
            sheet.auto_filter.ref = sheet.dimensions
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = f"A1:G{sheet.max_row}"
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.fitToPage = True
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.paperSize = sheet.PAPERSIZE_TABLOID
            sheet.page_margins.left = 0.4
            sheet.page_margins.right = 0.4
            sheet.page_margins.top = 0.5
            sheet.page_margins.bottom = 0.5
            sheet.oddHeader.left.text = SHEET_NAMES[raid["id"]]
            sheet.oddFooter.right.text = "Page &P of &N"
            sheet.print_title_rows = "1:1"
            sheet.sheet_view.showGridLines = False
            sheet.sheet_view.zoomScale = 110

    OUT.parent.mkdir(parents=True, exist_ok=True)
    book.save(OUT)
    print(f"wrote {OUT}")
    for sheet in book.worksheets:
        print(f"  {sheet.title}: {sheet.max_row - 1} spells")


if __name__ == "__main__":
    main()
