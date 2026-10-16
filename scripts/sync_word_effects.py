"""Put the website's spell effect text into the Word writeup."""

import json
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
DASH = re.compile(r"^(\s*[—–-]\s*)(.*)$", re.S)
ID_LINE = re.compile(r"^\s*(\d+)\s*[—–-]\s*")
TITLE_IDS = re.compile(r"\(([\d,\s]+)\)")


def element_text(element):
    return "".join(node.text or "" for node in element.iter(W + "t"))


def set_run_text(run, text):
    nodes = list(run.iter(W + "t"))
    if not nodes:
        node = OxmlElement("w:t")
        run.append(node)
        nodes = [node]
    nodes[0].text = text
    nodes[0].set(qn("xml:space"), "preserve")
    for extra in nodes[1:]:
        extra.text = ""


def insert_after(run, text):
    new_run = OxmlElement("w:r")
    properties = run.find(W + "rPr")
    if properties is not None:
        new_run.append(deepcopy(properties))
    node = OxmlElement("w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    new_run.append(node)
    run.addnext(new_run)


def sync_paragraph(paragraph, spells, effects):
    children = [child for child in paragraph._p if child.tag in (W + "r", W + "hyperlink")]
    for index, child in enumerate(children):
        if child.tag != W + "hyperlink":
            continue
        label = element_text(child).strip()
        if label.isdigit() and label in effects:
            following = child.getnext()
            if following is None or following.tag != W + "r" or not DASH.match(element_text(following)):
                continue
            current = DASH.match(element_text(following)).group(2).strip()
            if effects[label] and current != effects[label]:
                set_run_text(following, " — " + effects[label])
            continue
        matches = spells.get(label) or []
        if len(matches) != 1 or index + 3 >= len(children):
            continue
        effect = matches[0].get("effect") or ""
        if not effect or effect in element_text(paragraph._p):
            continue
        opening, spell_id, closing = children[index + 1 : index + 4]
        if element_text(opening).strip() != "(" or element_text(closing).strip() != ")":
            continue
        if element_text(spell_id).strip() != matches[0]["id"]:
            continue
        following = closing.getnext()
        if following is not None and following.tag == W + "r":
            dashed = DASH.match(element_text(following))
            if dashed:
                description = dashed.group(2).strip()
                if not description:
                    described = following.getnext()
                    if described is not None and described.tag == W + "r":
                        description = element_text(described).strip()
                if description and not description.startswith(effect):
                    set_run_text(following, " — " + effect + ". ")
                    if dashed.group(2).strip():
                        set_run_text(following, " — " + effect + ". " + description)
                continue
            # Leave mid-sentence mentions alone, e.g. "Spawns Name (id) every 270 seconds".
            if element_text(following).strip():
                continue
        insert_after(closing, " — " + effect)


def remove_duplicate_id_lines(document):
    """Drop `id — effect` lines that repeat a heading already showing Name (id) — effect."""
    removed = 0
    paragraphs = list(document.paragraphs)
    index = 1
    while index < len(paragraphs):
        current = paragraphs[index].text.strip()
        previous = paragraphs[index - 1].text.strip()
        match = ID_LINE.match(current)
        group = TITLE_IDS.search(previous)
        if match and group:
            title_ids = [part.strip() for part in group.group(1).split(",") if part.strip()]
            if title_ids == [match.group(1)]:
                parent = paragraphs[index]._p.getparent()
                if parent is not None:
                    parent.remove(paragraphs[index]._p)
                    removed += 1
                    paragraphs = list(document.paragraphs)
                    continue
        index += 1
    return removed


def main():
    data = json.loads((ROOT / "data" / "raids.json").read_text(encoding="utf-8"))
    effects = {
        match["id"]: match.get("effect") or ""
        for matches in data["spells"].values()
        for match in matches
    }
    document = Document(ROOT / "Info" / "FoF_Raid_Mechanics.docx")
    for paragraph in document.paragraphs:
        sync_paragraph(paragraph, data["spells"], effects)
    removed = remove_duplicate_id_lines(document)
    document.save(ROOT / "Info" / "FoF_Raid_Mechanics.docx")
    print(f"updated word writeup duplicates_removed={removed}")


if __name__ == "__main__":
    main()
